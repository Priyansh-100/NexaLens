import hashlib
import io
import mimetypes
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

import pdfplumber
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pydantic import BaseModel

from nexalens.core.config import get_settings
from nexalens.core.exceptions import DocumentProcessingError
from nexalens.core.logging import get_logger
from nexalens.models.schemas import Document, DocumentChunk
from nexalens.services.embeddings import embedding_service

logger = get_logger(__name__)
settings = get_settings()


class ProcessedDocument(BaseModel):
    document: Document
    chunks: list[DocumentChunk]


class DocumentProcessor:
    def __init__(self):
        self.chunk_size = settings.chunk_size
        self.chunk_overlap = settings.chunk_overlap
        self.max_size_bytes = settings.max_doc_size_mb * 1024 * 1024
        self.supported_extensions = set(settings.supported_extensions)

        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=self.chunk_size,
            chunk_overlap=self.chunk_overlap,
            length_function=len,
            separators=["\n\n", "\n", ". ", " ", ""],
        )

    def validate_file(self, filename: str, content: bytes) -> None:
        ext = Path(filename).suffix.lower()
        if ext not in self.supported_extensions:
            raise DocumentProcessingError(f"Unsupported file type: {ext}. Supported: {self.supported_extensions}")

        if len(content) > self.max_size_bytes:
            raise DocumentProcessingError(f"File too large: {len(content)} bytes (max: {self.max_size_bytes})")

    def extract_text(self, filename: str, content: bytes) -> str:
        ext = Path(filename).suffix.lower()

        try:
            if ext == ".pdf":
                return self._extract_pdf(content)
            elif ext in {".txt", ".md"}:
                return content.decode("utf-8")
            elif ext == ".csv":
                return self._extract_csv(content)
            elif ext in {".xlsx", ".xls"}:
                return self._extract_excel(content)
            elif ext == ".docx":
                return self._extract_docx(content)
            else:
                raise DocumentProcessingError(f"No extractor for: {ext}")
        except DocumentProcessingError:
            raise
        except Exception as e:
            logger.error("text_extraction_failed", filename=filename, error=str(e))
            raise DocumentProcessingError(f"Failed to extract text: {e}") from e

    def _extract_pdf(self, content: bytes) -> str:
        text_parts = []
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        return "\n\n".join(text_parts)

    def _extract_csv(self, content: bytes) -> str:
        import pandas as pd
        df = pd.read_csv(io.BytesIO(content))
        return df.to_markdown(index=False)

    def _extract_excel(self, content: bytes) -> str:
        import pandas as pd
        sheets = pd.read_excel(io.BytesIO(content), sheet_name=None)
        parts = []
        for name, df in sheets.items():
            parts.append(f"## Sheet: {name}\n{df.to_markdown(index=False)}")
        return "\n\n".join(parts)

    def _extract_docx(self, content: bytes) -> str:
        from docx import Document as DocxDocument
        doc = DocxDocument(io.BytesIO(content))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    def process(
        self,
        source_id: UUID,
        filename: str,
        content: bytes,
        metadata: dict[str, Any] | None = None,
    ) -> ProcessedDocument:
        self.validate_file(filename, content)
        text = self.extract_text(filename, content)

        if not text.strip():
            raise DocumentProcessingError("No text content extracted")

        doc_id = uuid4()
        doc_hash = hashlib.sha256(content).hexdigest()

        doc_metadata = metadata or {}
        doc_metadata.update(
            {
                "filename": filename,
                "content_hash": doc_hash,
                "size_bytes": len(content),
                "mime_type": mimetypes.guess_type(filename)[0],
            }
        )

        document = Document(
            id=doc_id,
            source_id=source_id,
            title=Path(filename).stem,
            content=text,
            metadata=doc_metadata,
        )

        chunks = self._create_chunks(document)
        document.chunk_count = len(chunks)

        return ProcessedDocument(document=document, chunks=chunks)

    def _create_chunks(self, document: Document) -> list[DocumentChunk]:
        texts = self.text_splitter.split_text(document.content)
        chunks = []

        for i, text in enumerate(texts):
            chunk = DocumentChunk(
                id=uuid4(),
                document_id=document.id,
                content=text,
                chunk_index=i,
                metadata={
                    **document.metadata,
                    "document_title": document.title,
                },
            )
            chunks.append(chunk)

        return chunks

    async def embed_chunks(self, chunks: list[DocumentChunk]) -> list[DocumentChunk]:
        texts = [chunk.content for chunk in chunks]
        embeddings = await embedding_service.embed_batch(texts)

        for chunk, embedding in zip(chunks, embeddings):
            chunk.embedding = embedding

        return chunks


document_processor = DocumentProcessor()