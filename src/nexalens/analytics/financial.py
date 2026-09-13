import base64
import io
from typing import Any

import numpy as np
import numpy_financial as npf
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, NumberFormat, PatternFill
from openpyxl.utils import get_column_letter

from nexalens.core.logging import get_logger
from nexalens.models.schemas import FinancialModelRequest, FinancialModelResult, FinancialModelType

logger = get_logger(__name__)


class FinancialModelingService:
    def __init__(self):
        self.default_assumptions = {
            FinancialModelType.DCF: {
                "revenue_growth": 0.15,
                "ebitda_margin": 0.25,
                "tax_rate": 0.21,
                "discount_rate": 0.12,
                "terminal_growth": 0.03,
                "capex_pct_revenue": 0.05,
                "depreciation_pct_revenue": 0.04,
                "working_capital_pct_revenue": 0.10,
            },
            FinancialModelType.NPV: {
                "discount_rate": 0.10,
            },
            FinancialModelType.IRR: {},
            FinancialModelType.PAYBACK: {},
            FinancialModelType.SENSITIVITY: {
                "base_npv": 0.0,
            },
        }

    def get_assumptions_template(self, model_type: FinancialModelType) -> dict[str, Any]:
        return self.default_assumptions.get(model_type, {}).copy()

    def calculate_npv(self, cash_flows: list[float], discount_rate: float) -> float:
        if not cash_flows:
            return 0.0
        return float(npf.npv(discount_rate, cash_flows))

    def calculate_irr(self, cash_flows: list[float]) -> float | None:
        if not cash_flows:
            return None
        try:
            irr = npf.irr(cash_flows)
            return float(irr) if irr is not None else None
        except Exception:
            return None

    def calculate_payback(self, cash_flows: list[float]) -> float | None:
        if not cash_flows:
            return None
        cumulative = 0.0
        for i, cf in enumerate(cash_flows):
            cumulative += cf
            if cumulative >= 0:
                if i == 0:
                    return 0.0
                prev_cumulative = cumulative - cf
                return i - 1 + abs(prev_cumulative) / abs(cf)
        return None

    def calculate_dcf(
        self,
        free_cash_flows: list[float],
        terminal_growth: float,
        discount_rate: float,
        tax_rate: float = 0.21,
    ) -> dict[str, float]:
        if not free_cash_flows:
            return {"dcf_value": 0.0, "terminal_value": 0.0, "pv_terminal": 0.0}

        n = len(free_cash_flows)
        pv_fcf = sum(fcf / (1 + discount_rate) ** (i + 1) for i, fcf in enumerate(free_cash_flows))

        terminal_fcf = free_cash_flows[-1] * (1 + terminal_growth)
        terminal_value = terminal_fcf / (discount_rate - terminal_growth)
        pv_terminal = terminal_value / (1 + discount_rate) ** n

        return {
            "dcf_value": pv_fcf + pv_terminal,
            "terminal_value": terminal_value,
            "pv_terminal": pv_terminal,
            "pv_fcf": pv_fcf,
        }

    def build_sensitivity_table(
        self,
        base_assumptions: dict[str, Any],
        variables: list[str],
        ranges: dict[str, tuple[float, float]],
        steps: int = 5,
    ) -> dict[str, list[float]]:
        table = {}
        for var in variables:
            if var not in ranges:
                continue
            min_val, max_val = ranges[var]
            vals = np.linspace(min_val, max_val, steps)
            npvs = []
            for val in vals:
                assumptions = base_assumptions.copy()
                assumptions[var] = val
                if "cash_flows" in assumptions:
                    npv = self.calculate_npv(assumptions["cash_flows"], assumptions.get("discount_rate", 0.1))
                    npvs.append(npv)
            table[var] = npvs
        return table

    def export_to_excel(self, result: FinancialModelResult) -> str:
        wb = Workbook()
        ws = wb.active
        ws.title = "Financial Model"

        header_font = Font(bold=True, size=12)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font_white = Font(bold=True, size=12, color="FFFFFF")
        currency_format = "#,##0.00"
        pct_format = "0.00%"

        row = 1
        ws.cell(row=row, column=1, value=f"{result.model_type.value.upper()} Model Results").font = Font(bold=True, size=14)
        row += 2

        ws.cell(row=row, column=1, value="Assumptions Used").font = header_font
        ws.cell(row=row, column=1).fill = header_fill
        ws.cell(row=row, column=1).font = header_font_white
        row += 1

        for key, value in result.assumptions_used.items():
            ws.cell(row=row, column=1, value=key.replace("_", " ").title())
            ws.cell(row=row, column=2, value=value)
            if isinstance(value, float) and abs(value) < 1:
                ws.cell(row=row, column=2).number_format = pct_format
            elif isinstance(value, float):
                ws.cell(row=row, column=2).number_format = currency_format
            row += 1

        row += 1
        ws.cell(row=row, column=1, value="Results").font = header_font
        ws.cell(row=row, column=1).fill = header_fill
        ws.cell(row=row, column=1).font = header_font_white
        row += 1

        if result.npv is not None:
            ws.cell(row=row, column=1, value="NPV")
            ws.cell(row=row, column=2, value=result.npv).number_format = currency_format
            row += 1
        if result.irr is not None:
            ws.cell(row=row, column=1, value="IRR")
            ws.cell(row=row, column=2, value=result.irr).number_format = pct_format
            row += 1
        if result.payback_period is not None:
            ws.cell(row=row, column=1, value="Payback Period (years)")
            ws.cell(row=row, column=2, value=result.payback_period).number_format = "0.00"
            row += 1
        if result.dcf_value is not None:
            ws.cell(row=row, column=1, value="DCF Value")
            ws.cell(row=row, column=2, value=result.dcf_value).number_format = currency_format
            row += 1

        if result.sensitivity_table:
            row += 1
            ws.cell(row=row, column=1, value="Sensitivity Analysis").font = header_font
            ws.cell(row=row, column=1).fill = header_fill
            ws.cell(row=row, column=1).font = header_font_white
            row += 1

            for var, npvs in result.sensitivity_table.items():
                ws.cell(row=row, column=1, value=f"{var} Sensitivity")
                row += 1
                for npv in npvs:
                    ws.cell(row=row, column=1, value="")
                    ws.cell(row=row, column=2, value=npv).number_format = currency_format
                    row += 1

        for col in range(1, 3):
            ws.column_dimensions[get_column_letter(col)].width = 30

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return base64.b64encode(buffer.read()).decode("utf-8")

    def run_model(self, request: FinancialModelRequest) -> FinancialModelResult:
        assumptions = {**self.get_assumptions_template(request.model_type), **request.assumptions}
        cash_flows = request.cash_flows or []

        result = FinancialModelResult(
            model_type=request.model_type,
            assumptions_used=assumptions,
            explanation="",
        )

        if request.model_type == FinancialModelType.NPV:
            discount_rate = assumptions.get("discount_rate", 0.1)
            result.npv = self.calculate_npv(cash_flows, discount_rate)

        elif request.model_type == FinancialModelType.IRR:
            result.irr = self.calculate_irr(cash_flows)
            if result.irr is not None:
                result.npv = self.calculate_npv(cash_flows, result.irr)

        elif request.model_type == FinancialModelType.PAYBACK:
            result.payback_period = self.calculate_payback(cash_flows)

        elif request.model_type == FinancialModelType.DCF:
            discount_rate = assumptions.get("discount_rate", 0.12)
            terminal_growth = assumptions.get("terminal_growth", 0.03)
            dcf_result = self.calculate_dcf(cash_flows, terminal_growth, discount_rate)
            result.dcf_value = dcf_result["dcf_value"]
            result.npv = dcf_result["dcf_value"]

        elif request.model_type == FinancialModelType.SENSITIVITY:
            variables = assumptions.get("sensitivity_variables", [])
            ranges = assumptions.get("sensitivity_ranges", {})
            if variables and ranges:
                result.sensitivity_table = self.build_sensitivity_table(
                    assumptions, variables, ranges
                )
                base_discount = assumptions.get("discount_rate", 0.1)
                result.npv = self.calculate_npv(cash_flows, base_discount)

        if request.output_format == "excel":
            result.excel_base64 = self.export_to_excel(result)

        result.explanation = self._generate_explanation(result, assumptions)
        return result

    def _generate_explanation(self, result: FinancialModelResult, assumptions: dict) -> str:
        parts = [f"{result.model_type.value.upper()} Analysis Complete.\n"]

        if result.model_type == FinancialModelType.NPV and result.npv is not None:
            parts.append(f"NPV: ${result.npv:,.2f} at {assumptions.get('discount_rate', 0.1):.1%} discount rate.")
            parts.append("Positive NPV indicates value creation." if result.npv > 0 else "Negative NPV indicates value destruction.")

        elif result.model_type == FinancialModelType.IRR and result.irr is not None:
            parts.append(f"IRR: {result.irr:.2%}.")
            dr = assumptions.get("discount_rate", 0.1)
            parts.append(f"Project exceeds hurdle rate of {dr:.1%}." if result.irr > dr else f"Project below hurdle rate of {dr:.1%}.")

        elif result.model_type == FinancialModelType.DCF and result.dcf_value is not None:
            parts.append(f"DCF Enterprise Value: ${result.dcf_value:,.2f}.")
            parts.append(f"Terminal growth assumption: {assumptions.get('terminal_growth', 0.03):.1%}.")

        elif result.model_type == FinancialModelType.PAYBACK and result.payback_period is not None:
            parts.append(f"Payback Period: {result.payback_period:.2f} years.")

        elif result.model_type == FinancialModelType.SENSITIVITY and result.sensitivity_table:
            parts.append("Sensitivity analysis shows NPV impact of key variables.")
            for var, npvs in result.sensitivity_table.items():
                parts.append(f"{var}: NPV range ${min(npvs):,.0f} to ${max(npvs):,.0f}.")

        return "\n".join(parts)


financial_modeling_service = FinancialModelingService()