import asyncio
import base64
import os
import smtplib
from datetime import datetime
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

import aiofiles
import httpx
import jinja2
import pandas as pd
from apscheduler.schedulers.asyncio import AsyncIOScheduler
from apscheduler.triggers.cron import CronTrigger
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from nexalens.core.config import get_settings
from nexalens.core.logging import get_logger
from nexalens.models.database import ReportExecutionModel, ReportScheduleModel
from nexalens.models.schemas import (
    ReportExecution,
    ReportFormat,
    ReportSchedule,
    ReportScheduleCreate,
    ReportScheduleUpdate,
)
from nexalens.rag.orchestrator import rag_orchestrator
from nexalens.services.llm import llm_service

logger = get_logger(__name__)
settings = get_settings()


class ReportSchedulerService:
    def __init__(self):
        self.scheduler = AsyncIOScheduler()
        self.templates_dir = Path("templates/reports")
        self.templates_dir.mkdir(parents=True, exist_ok=True)
        self.jinja_env = jinja2.Environment(
            loader=jinja2.FileSystemLoader(str(self.templates_dir)),
            autoescape=jinja2.select_autoescape(["html", "xml"]),
        )
        self.jinja_env.filters["format_currency"] = self._format_currency
        self.jinja_env.filters["format_pct"] = self._format_pct
        self.jinja_env.filters["format_number"] = self._format_number
        self._job_map: dict[UUID, Any] = {}

    def _format_currency(self, value: float) -> str:
        if value is None:
            return "N/A"
        if abs(value) >= 1e9:
            return f"${value/1e9:.2f}B"
        elif abs(value) >= 1e6:
            return f"${value/1e6:.2f}M"
        elif abs(value) >= 1e3:
            return f"${value/1e3:.2f}K"
        return f"${value:,.2f}"

    def _format_pct(self, value: float) -> str:
        if value is None:
            return "N/A"
        return f"{value:.2%}"

    def _format_number(self, value: float) -> str:
        if value is None:
            return "N/A"
        if abs(value) >= 1e6:
            return f"{value/1e6:.2f}M"
        elif abs(value) >= 1e3:
            return f"{value/1e3:.2f}K"
        return f"{value:,.2f}"

    async def create_schedule(
        self,
        schedule_in: ReportScheduleCreate,
        session: AsyncSession,
        user_id: UUID,
    ) -> ReportSchedule:
        schedule = ReportScheduleModel(
            name=schedule_in.name,
            query=schedule_in.query,
            data_source_ids=schedule_in.data_source_ids,
            cron_expression=schedule_in.cron_expression,
            recipients=schedule_in.recipients,
            format=schedule_in.format.value,
            template=schedule_in.template,
            created_by=user_id,
        )
        session.add(schedule)
        await session.flush()
        await session.refresh(schedule)

        if schedule.is_active:
            await self._schedule_job(schedule)

        logger.info("report_schedule_created", schedule_id=str(schedule.id), name=schedule.name)
        return self._to_schema(schedule)

    async def update_schedule(
        self,
        schedule_id: UUID,
        updates: ReportScheduleUpdate,
        session: AsyncSession,
    ) -> ReportSchedule | None:
        schedule = await session.get(ReportScheduleModel, schedule_id)
        if not schedule:
            return None

        update_data = updates.model_dump(exclude_unset=True)
        for field, value in update_data.items():
            if field == "format" and value:
                setattr(schedule, field, value.value)
            else:
                setattr(schedule, field, value)

        await session.flush()

        await self._remove_job(schedule_id)
        if schedule.is_active:
            await self._schedule_job(schedule)

        logger.info("report_schedule_updated", schedule_id=str(schedule_id))
        return self._to_schema(schedule)

    async def delete_schedule(self, schedule_id: UUID, session: AsyncSession) -> bool:
        schedule = await session.get(ReportScheduleModel, schedule_id)
        if not schedule:
            return False

        await self._remove_job(schedule_id)
        await session.delete(schedule)
        logger.info("report_schedule_deleted", schedule_id=str(schedule_id))
        return True

    async def get_schedules(
        self,
        session: AsyncSession,
        user_id: UUID | None = None,
    ) -> list[ReportSchedule]:
        stmt = select(ReportScheduleModel)
        if user_id:
            stmt = stmt.where(ReportScheduleModel.created_by == user_id)
        result = await session.execute(stmt)
        return [self._to_schema(s) for s in result.scalars().all()]

    async def get_schedule(self, schedule_id: UUID, session: AsyncSession) -> ReportSchedule | None:
        schedule = await session.get(ReportScheduleModel, schedule_id)
        return self._to_schema(schedule) if schedule else None

    async def run_now(self, schedule_id: UUID, session: AsyncSession) -> ReportExecution:
        schedule = await session.get(ReportScheduleModel, schedule_id)
        if not schedule:
            raise ValueError("Schedule not found")
        return await self._execute_schedule(schedule, session, manual=True)

    async def get_executions(
        self,
        session: AsyncSession,
        schedule_id: UUID | None = None,
        limit: int = 50,
        user_id: UUID | None = None,
    ) -> list[ReportExecution]:
        stmt = (
            select(ReportExecutionModel)
            .join(ReportScheduleModel, ReportExecutionModel.schedule_id == ReportScheduleModel.id)
            .order_by(ReportExecutionModel.started_at.desc())
            .limit(limit)
        )
        if schedule_id:
            stmt = stmt.where(ReportExecutionModel.schedule_id == schedule_id)
        if user_id:
            stmt = stmt.where(ReportScheduleModel.created_by == user_id)
        result = await session.execute(stmt)
        return [self._to_execution_schema(e) for e in result.scalars().all()]

    async def _schedule_job(self, schedule: ReportScheduleModel) -> None:
        try:
            trigger = CronTrigger.from_crontab(schedule.cron_expression)
            job = self.scheduler.add_job(
                self._run_scheduled_job,
                trigger=trigger,
                args=[schedule.id],
                id=f"report_{schedule.id}",
                replace_existing=True,
            )
            self._job_map[schedule.id] = job

            next_run = job.next_run_time
            if next_run:
                schedule.next_run = next_run.replace(tzinfo=None)

            logger.info("report_job_scheduled", schedule_id=str(schedule.id), next_run=str(next_run))
        except Exception as e:
            logger.error("schedule_job_failed", schedule_id=str(schedule.id), error=str(e))

    async def _remove_job(self, schedule_id: UUID) -> None:
        if schedule_id in self._job_map:
            try:
                self.scheduler.remove_job(f"report_{schedule_id}")
                del self._job_map[schedule_id]
            except Exception as e:
                logger.warning("remove_job_failed", schedule_id=str(schedule_id), error=str(e))

    async def _run_scheduled_job(self, schedule_id: UUID) -> None:
        from nexalens.models.session import async_session_maker

        async with async_session_maker() as session:
            schedule = await session.get(ReportScheduleModel, schedule_id)
            if schedule and schedule.is_active:
                await self._execute_schedule(schedule, session, manual=False)

    async def _execute_schedule(
        self,
        schedule: ReportScheduleModel,
        session: AsyncSession,
        manual: bool = False,
    ) -> ReportExecution:
        execution = ReportExecutionModel(
            schedule_id=schedule.id,
            status="running",
        )
        session.add(execution)
        await session.flush()

        start_time = datetime.utcnow()

        try:
            from nexalens.models.schemas import QueryRequest, QueryIntent

            request = QueryRequest(
                question=schedule.query,
                data_source_ids=schedule.data_source_ids,
                intent=QueryIntent.HYBRID,
                max_results=100,
                include_sql=True,
                include_sources=True,
            )

            response = await rag_orchestrator.process_query(request, session)

            content, filename = await self._render_report(response, schedule)

            for recipient in schedule.recipients:
                await self._deliver_report(recipient, content, filename, schedule.format)

            execution.status = "success"
            execution.output_path = filename
            schedule.last_run = datetime.utcnow()

            logger.info("report_executed_success", schedule_id=str(schedule.id), execution_id=str(execution.id))

        except Exception as e:
            execution.status = "failed"
            execution.error = str(e)
            logger.error("report_execution_failed", schedule_id=str(schedule.id), error=str(e))

        finally:
            execution.completed_at = datetime.utcnow()
            await session.flush()

        return self._to_execution_schema(execution)

    async def _render_report(
        self,
        response,
        schedule: ReportScheduleModel,
    ) -> tuple[bytes, str]:
        template_name = schedule.template or "default"
        template = self._get_template(template_name, schedule.format)

        context = {
            "schedule_name": schedule.name,
            "generated_at": datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC"),
            "question": response.question,
            "answer": response.answer,
            "sql_result": response.sql_result,
            "doc_results": response.document_results,
            "confidence": response.confidence,
        }

        if schedule.format == ReportFormat.HTML:
            content = template.render(**context).encode("utf-8")
            filename = f"report_{schedule.id}_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.html"
        elif schedule.format == ReportFormat.CSV:
            content = self._render_csv(response).encode("utf-8")
            filename = f"report_{schedule.id}_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.csv"
        elif schedule.format == ReportFormat.EXCEL:
            content = self._render_excel(response)
            filename = f"report_{schedule.id}_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
        else:
            html_content = template.render(**context)
            content = await self._html_to_pdf(html_content)
            filename = f"report_{schedule.id}_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.pdf"

        return content, filename

    def _get_template(self, name: str, fmt: ReportFormat) -> jinja2.Template:
        template_file = f"{name}.j2"
        if fmt == ReportFormat.HTML:
            template_file = f"{name}.html.j2"

        try:
            return self.jinja_env.get_template(template_file)
        except jinja2.exceptions.TemplateNotFound:
            return self._default_template(fmt)

    def _default_template(self, fmt: ReportFormat) -> jinja2.Template:
        if fmt == ReportFormat.HTML:
            return self.jinja_env.from_string(DEFAULT_HTML_TEMPLATE)
        return self.jinja_env.from_string(DEFAULT_TEXT_TEMPLATE)

    def _render_csv(self, response) -> str:
        lines = []
        if response.sql_result and response.sql_result.rows:
            df = pd.DataFrame(response.sql_result.rows)
            lines.append("# SQL Results")
            lines.append(df.to_csv(index=False))
        return "\n".join(lines)

    def _render_excel(self, response) -> bytes:
        import openpyxl
        from openpyxl.styles import Font

        wb = openpyxl.Workbook()

        if response.sql_result and response.sql_result.rows:
            ws1 = wb.active
            ws1.title = "SQL Results"
            for col_idx, col_name in enumerate(response.sql_result.columns, 1):
                cell = ws1.cell(row=1, column=col_idx, value=col_name)
                cell.font = Font(bold=True)
            for row_idx, row in enumerate(response.sql_result.rows, 2):
                for col_idx, col_name in enumerate(response.sql_result.columns, 1):
                    ws1.cell(row=row_idx, column=col_idx, value=row.get(col_name))

        ws2 = wb.create_sheet("Summary")
        ws2.cell(row=1, column=1, value="Question").font = Font(bold=True)
        ws2.cell(row=1, column=2, value=response.question)
        ws2.cell(row=2, column=1, value="Answer").font = Font(bold=True)
        ws2.cell(row=2, column=2, value=response.answer)
        ws2.cell(row=3, column=1, value="Confidence").font = Font(bold=True)
        ws2.cell(row=3, column=2, value=f"{response.confidence:.1%}")

        import io
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.read()

    async def _html_to_pdf(self, html: str) -> bytes:
        try:
            from weasyprint import HTML
            return HTML(string=html).write_pdf()
        except Exception as e:
            logger.warning("weasyprint_failed_fallback", error=str(e))
            return html.encode("utf-8")

    async def _deliver_report(
        self,
        recipient: str,
        content: bytes,
        filename: str,
        fmt: ReportFormat,
    ) -> None:
        if recipient.startswith("http"):
            await self._deliver_webhook(recipient, content, filename)
        elif "@" in recipient:
            await self._deliver_email(recipient, content, filename, fmt)
        else:
            logger.warning("unknown_recipient_type", recipient=recipient)

    async def _deliver_email(
        self,
        email: str,
        content: bytes,
        filename: str,
        fmt: ReportFormat,
    ) -> None:
        smtp_host = os.getenv("SMTP_HOST")
        smtp_port = int(os.getenv("SMTP_PORT", "587"))
        smtp_user = os.getenv("SMTP_USER")
        smtp_pass = os.getenv("SMTP_PASS")

        if not all([smtp_host, smtp_user, smtp_pass]):
            logger.warning("email_not_configured", recipient=email)
            return

        msg = MIMEMultipart()
        msg["Subject"] = f"NexaLens Report: {filename}"
        msg["From"] = smtp_user
        msg["To"] = email

        msg.attach(MIMEText(f"Your scheduled report is attached.\n\nFormat: {fmt.value}\nGenerated: {datetime.utcnow()}", "plain"))

        attachment = MIMEApplication(content, Name=filename)
        attachment["Content-Disposition"] = f'attachment; filename="{filename}"'
        msg.attach(attachment)

        try:
            with smtplib.SMTP(smtp_host, smtp_port) as server:
                server.starttls()
                server.login(smtp_user, smtp_pass)
                server.send_message(msg)
            logger.info("report_email_sent", recipient=email)
        except Exception as e:
            logger.error("email_send_failed", recipient=email, error=str(e))

    async def _deliver_webhook(self, url: str, content: bytes, filename: str) -> None:
        try:
            async with httpx.AsyncClient(timeout=30.0) as client:
                files = {"file": (filename, content)}
                await client.post(url, files=files)
            logger.info("report_webhook_delivered", url=url)
        except Exception as e:
            logger.error("webhook_delivery_failed", url=url, error=str(e))

    def start(self) -> None:
        self.scheduler.start()
        logger.info("report_scheduler_started")

    def shutdown(self) -> None:
        self.scheduler.shutdown()
        logger.info("report_scheduler_shutdown")

    def _to_schema(self, model: ReportScheduleModel) -> ReportSchedule:
        return ReportSchedule(
            id=model.id,
            name=model.name,
            query=model.query,
            data_source_ids=model.data_source_ids,
            cron_expression=model.cron_expression,
            recipients=model.recipients,
            format=ReportFormat(model.format),
            template=model.template,
            is_active=model.is_active,
            created_by=model.created_by,
            created_at=model.created_at,
            last_run=model.last_run,
            next_run=model.next_run,
        )

    def _to_execution_schema(self, model: ReportExecutionModel) -> ReportExecution:
        return ReportExecution(
            id=model.id,
            schedule_id=model.schedule_id,
            status=model.status,
            output_path=model.output_path,
            error=model.error,
            started_at=model.started_at,
            completed_at=model.completed_at,
        )


report_scheduler = ReportSchedulerService()


DEFAULT_HTML_TEMPLATE = """
<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>{{ schedule_name }}</title>
    <style>
        body { font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; margin: 40px; line-height: 1.6; }
        h1 { color: #1a1a2e; border-bottom: 2px solid #4472C4; padding-bottom: 10px; }
        h2 { color: #2d2d44; margin-top: 30px; }
        .meta { color: #666; font-size: 0.9em; margin-bottom: 20px; }
        .section { margin: 20px 0; }
        .sql-table { border-collapse: collapse; width: 100%; margin: 15px 0; }
        .sql-table th, .sql-table td { border: 1px solid #ddd; padding: 8px; text-align: left; }
        .sql-table th { background: #4472C4; color: white; }
        .sql-table tr:nth-child(even) { background: #f9f9f9; }
        .doc-chunk { background: #f5f5f5; padding: 15px; margin: 10px 0; border-left: 4px solid #4472C4; }
        .confidence { display: inline-block; padding: 5px 15px; background: #e8f5e9; color: #2e7d32; border-radius: 20px; font-weight: bold; }
    </style>
</head>
<body>
    <h1>{{ schedule_name }}</h1>
    <div class="meta">Generated: {{ generated_at }} | Question: {{ question }}</div>

    <div class="section">
        <h2>Answer</h2>
        <p>{{ answer }}</p>
    </div>

    {% if sql_result and sql_result.rows %}
    <div class="section">
        <h2>SQL Results ({{ sql_result.row_count }} rows)</h2>
        <table class="sql-table">
            <thead>
                <tr>{% for col in sql_result.columns %}<th>{{ col }}</th>{% endfor %}</tr>
            </thead>
            <tbody>
                {% for row in sql_result.rows[:50] %}
                <tr>{% for col in sql_result.columns %}<td>{{ row.get(col, '') }}</td>{% endfor %}</tr>
                {% endfor %}
            </tbody>
        </table>
        {% if sql_result.row_count > 50 %}<p><em>Showing first 50 of {{ sql_result.row_count }} rows</em></p>{% endif %}
    </div>
    {% endif %}

    {% if doc_results %}
    <div class="section">
        <h2>Document Sources</h2>
        {% for doc in doc_results %}
        <div class="doc-chunk">
            <strong>{{ doc.title }}</strong> (relevance: {{ "%.0f"|format(doc.score * 100) }}%)<br>
            {{ doc.chunk_content[:300] }}...
        </div>
        {% endfor %}
    </div>
    {% endif %}

    <div class="section">
        <span class="confidence">Confidence: {{ "%.0f"|format(confidence * 100) }}%</span>
    </div>
</body>
</html>
"""

DEFAULT_TEXT_TEMPLATE = """
{{ schedule_name }}
Generated: {{ generated_at }}
Question: {{ question }}

=== ANSWER ===
{{ answer }}

{% if sql_result and sql_result.rows %}
=== SQL RESULTS ({{ sql_result.row_count }} rows) ===
Columns: {{ sql_result.columns | join(', ') }}
{{ sql_result.rows[:20] }}
{% endif %}

{% if doc_results %}
=== DOCUMENT SOURCES ===
{% for doc in doc_results %}
- {{ doc.title }} ({{ "%.0f"|format(doc.score * 100) }}%): {{ doc.chunk_content[:200] }}...
{% endfor %}
{% endif %}

Confidence: {{ "%.0f"|format(confidence * 100) }}%
"""